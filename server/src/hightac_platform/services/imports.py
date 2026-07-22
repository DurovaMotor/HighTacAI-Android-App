from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from hightac_platform.auth.context import Actor
from hightac_platform.db.models import ImportJob, ImportJobRow, Product
from hightac_platform.domain.enums import ImportJobStatus, ProductSource
from hightac_platform.domain.errors import ConflictError, NotFoundError, ValidationError
from hightac_platform.services.audit import append_operation_log
from hightac_platform.utils import normalize_product_code, utc_ms


MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000
CODE_HEADERS = {"产品编码", "产品编码*", "product code", "product_code"}
NAME_HEADERS = {"产品名称", "product name", "product_name"}


class ProductImportService:
    def template(self) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Products"
        worksheet.append(["产品编码*", "产品名称"])
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        worksheet.column_dimensions["A"].width = 28
        worksheet.column_dimensions["B"].width = 36
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def preview(
        self, session: Session, actor: Actor, *, filename: str, content: bytes
    ) -> ImportJob:
        if not content or len(content) > MAX_IMPORT_BYTES:
            raise ValidationError("Excel file is empty or exceeds the 5 MB limit.")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError("Uploaded file is not a readable Excel workbook.") from exc
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValidationError("Excel workbook has no header row.")
        normalized_headers = [str(value or "").strip().lower() for value in headers]
        code_index = _find_header(normalized_headers, CODE_HEADERS)
        name_index = _find_header(normalized_headers, NAME_HEADERS, required=False)

        job = ImportJob(
            filename=(filename or "products.xlsx")[:255],
            actor_type=actor.actor_type.value,
            actor_id=actor.actor_id,
        )
        session.add(job)
        session.flush()
        seen_codes: set[str] = set()
        total = valid = errors = 0
        for row_number, values in enumerate(rows, start=2):
            if row_number > MAX_IMPORT_ROWS + 1:
                raise ValidationError(f"Excel import is limited to {MAX_IMPORT_ROWS} data rows.")
            if not any(value is not None and str(value).strip() for value in values):
                continue
            total += 1
            raw_code = values[code_index] if code_index < len(values) else None
            raw_name = values[name_index] if name_index is not None and name_index < len(values) else None
            error_message = None
            normalized_code = None
            try:
                normalized_code = normalize_product_code(str(raw_code or ""))
                if normalized_code in seen_codes:
                    error_message = "Duplicate product code in workbook."
                seen_codes.add(normalized_code)
            except ValueError as exc:
                error_message = str(exc)
            product_name = str(raw_name).strip()[:256] if raw_name is not None else None
            if error_message:
                errors += 1
            else:
                valid += 1
            session.add(
                ImportJobRow(
                    job_id=job.id,
                    row_number=row_number,
                    product_code=normalized_code,
                    product_name=product_name or None,
                    is_valid=error_message is None,
                    error_message=error_message,
                )
            )
        if total == 0:
            raise ValidationError("Excel workbook contains no product rows.")
        job.row_count = total
        job.valid_count = valid
        job.error_count = errors
        job.status = (
            ImportJobStatus.INVALID.value if errors else ImportJobStatus.PREVIEWED.value
        )
        append_operation_log(
            session,
            event_type="product_import.previewed",
            actor=actor,
            result_summary={"job_id": job.id, "rows": total, "errors": errors},
        )
        session.commit()
        return job

    def commit(self, session: Session, actor: Actor, job_id: str) -> ImportJob:
        job = session.get(ImportJob, job_id)
        if job is None:
            raise NotFoundError("Import job was not found.")
        if job.status == ImportJobStatus.COMMITTED.value:
            return job
        if job.status != ImportJobStatus.PREVIEWED.value or job.error_count:
            raise ConflictError("Import job contains validation errors and cannot be committed.")
        rows = list(
            session.scalars(
                select(ImportJobRow)
                .where(ImportJobRow.job_id == job.id)
                .order_by(ImportJobRow.row_number)
            )
        )
        existing = {
            product.product_code: product
            for product in session.scalars(
                select(Product).where(
                    Product.product_code.in_([row.product_code for row in rows if row.product_code])
                )
            )
        }
        now = utc_ms()
        for row in rows:
            if not row.product_code or not row.is_valid:
                raise ConflictError("Import job changed after preview and cannot be committed.")
            product = existing.get(row.product_code)
            if product is None:
                product = Product(
                    product_code=row.product_code,
                    product_name=row.product_name,
                    source=ProductSource.EXCEL.value,
                )
                session.add(product)
                existing[row.product_code] = product
            elif row.product_name:
                product.product_name = row.product_name
                product.updated_at_ms = now
        job.status = ImportJobStatus.COMMITTED.value
        job.committed_at_ms = now
        append_operation_log(
            session,
            event_type="product_import.committed",
            actor=actor,
            result_summary={"job_id": job.id, "rows": job.valid_count},
        )
        session.commit()
        return job


def _find_header(headers: list[str], candidates: set[str], required: bool = True) -> int | None:
    lowered = {candidate.lower() for candidate in candidates}
    for index, header in enumerate(headers):
        if header in lowered:
            return index
    if required:
        raise ValidationError("Excel header must include 产品编码*.")
    return None
