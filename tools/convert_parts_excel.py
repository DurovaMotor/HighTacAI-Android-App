#!/usr/bin/env python3
"""Convert HighTac motorcycle parts Excel data into Android JSONL assets."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


DEFAULT_INPUT = Path(r"D:\CodeProject\logo\AllMotorPartsImformation.xlsx")
DEFAULT_OUTPUT_DIR = Path("app/src/main/assets")

FIELD_MAP = {
    "产品编码": "code",
    "产品中文名称": "nameCn",
    "产品英文名称": "nameEn",
    "品牌": "brand",
    "通用车型": "models",
    "规格型号": "spec",
    "采购人员": "buyer",
    "单位（中文）": "unit",
    "销售单价/元（标准价）": "standardPrice",
    "销售单价/元（最新价）": "latestPrice",
    "售价最新日期": "latestPriceDate",
    "售价最新数量": "latestPriceQty",
    "每箱数量(CTN)": "ctnQty",
    "毛重(KG)": "grossWeightKg",
    "长(CM)": "lengthCm",
    "宽(CM)": "widthCm",
    "高(CM)": "heightCm",
    "体积(CBM)": "volumeCbm",
    "备注": "remark",
}

TEXT_FIELDS = {
    "code",
    "nameCn",
    "nameEn",
    "brand",
    "models",
    "spec",
    "buyer",
    "unit",
    "remark",
}

NUMBER_FIELDS = {
    "standardPrice",
    "latestPrice",
    "latestPriceQty",
    "ctnQty",
    "grossWeightKg",
    "lengthCm",
    "widthCm",
    "heightCm",
    "volumeCbm",
}

OUTPUT_FIELDS = [
    "code",
    "nameCn",
    "nameEn",
    "brand",
    "models",
    "spec",
    "buyer",
    "unit",
    "standardPrice",
    "latestPrice",
    "latestPriceDate",
    "latestPriceQty",
    "ctnQty",
    "grossWeightKg",
    "lengthCm",
    "widthCm",
    "heightCm",
    "volumeCbm",
    "remark",
    "searchText",
]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def clean_number(value: Any) -> int | float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    text = clean_text(value)
    if not text:
        return None
    text = text.replace(",", "")
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def clean_date(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean_text(value)
    if not text:
        return None

    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def code_variants(code: str) -> list[str]:
    compact = re.sub(r"[\s\-_]+", "", code)
    spaced = re.sub(r"[\s\-_]+", " ", code).strip()
    variants = [code, spaced, compact]
    return [value for index, value in enumerate(variants) if value and value not in variants[:index]]


def build_search_text(record: dict[str, Any]) -> str:
    parts: list[str] = []
    for value in code_variants(record.get("code", "")):
        parts.append(value)

    for field in ("nameCn", "nameEn", "brand", "models", "spec", "remark"):
        value = clean_text(record.get(field))
        if value:
            parts.append(value)

    return re.sub(r"\s+", " ", " ".join(parts)).strip()


def empty_record() -> dict[str, Any]:
    record: dict[str, Any] = {}
    for field in OUTPUT_FIELDS:
        if field in NUMBER_FIELDS or field == "latestPriceDate":
            record[field] = None
        else:
            record[field] = ""
    return record


def convert(input_path: Path, output_dir: Path) -> int:
    if not input_path.exists():
        raise FileNotFoundError(f"Excel source not found: {input_path}")

    workbook = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    mapped_headers = [FIELD_MAP.get(clean_text(header)) for header in headers]

    records: list[dict[str, Any]] = []
    for row in rows:
        record = empty_record()
        for index, cell_value in enumerate(row):
            if index >= len(mapped_headers):
                continue
            field = mapped_headers[index]
            if not field:
                continue
            if field in TEXT_FIELDS:
                record[field] = clean_text(cell_value)
            elif field in NUMBER_FIELDS:
                record[field] = clean_number(cell_value)
            elif field == "latestPriceDate":
                record[field] = clean_date(cell_value)

        if not any(record.get(field) for field in OUTPUT_FIELDS if field != "searchText"):
            continue
        record["searchText"] = build_search_text(record)
        records.append({field: record[field] for field in OUTPUT_FIELDS})

    output_dir.mkdir(parents=True, exist_ok=True)
    jsonl_path = output_dir / "parts_catalog.jsonl"
    meta_path = output_dir / "parts_catalog_meta.json"

    with jsonl_path.open("w", encoding="utf-8", newline="\n") as file:
        for record in records:
            file.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")))
            file.write("\n")

    meta = {
        "source": input_path.name,
        "totalRows": len(records),
        "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
        "version": 1,
    }
    meta_path.write_text(
        json.dumps(meta, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return len(records)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert company motorcycle parts Excel file to Android JSONL assets."
    )
    parser.add_argument(
        "input",
        nargs="?",
        type=Path,
        default=DEFAULT_INPUT,
        help=f"Excel source path. Default: {DEFAULT_INPUT}",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Android assets output directory. Default: {DEFAULT_OUTPUT_DIR}",
    )
    args = parser.parse_args()

    total = convert(args.input, args.output_dir)
    print(f"Generated {total} records in {args.output_dir}")


if __name__ == "__main__":
    main()
